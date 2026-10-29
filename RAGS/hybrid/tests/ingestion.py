# load the sources from sample_files
# parse the data
# create the chunks
import io
import re
from pathlib import Path
import pandas as pd
import fitz
from docx import Document

def load_file(filename: str, file_bytes: bytes):
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        text = file_bytes.decode("utf-8", errors="replace")
    elif suffix == ".pdf":
        pages = []
        with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
            for n, page in enumerate(pdf, 1):
                pages.append(f"[Page {n}]\n{page.get_text()}")
        text = "\n\n".join(pages)
    elif suffix == ".docx":
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    elif suffix == ".csv":
        df = pd.read_csv(io.BytesIO(file_bytes)).fillna("")
        text = "\n".join(
            f"Row {i+1}: " + "; ".join(f"{k}: {v}" for k, v in row.items())
            for i, row in enumerate(df.to_dict(orient="records"))
        )
    elif suffix == ".xlsx":
        parts = []
        for sheet, df in pd.read_excel(io.BytesIO(file_bytes), sheet_name=None).items():
            parts.append(f"Sheet: {sheet}")
            df = df.fillna("")
            parts.extend(
                f"Row {i+1}: " + "; ".join(f"{k}: {v}" for k, v in row.items())
                for i, row in enumerate(df.to_dict(orient="records"))
            )
        text = "\n".join(parts)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    return text.strip()

def chunk_text(text: str, chunk_size: int = 900, overlap: int = 150):
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size")
    chunks, start = [], 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
            if boundary > start + int(chunk_size * 0.55):
                end = boundary + 1
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks